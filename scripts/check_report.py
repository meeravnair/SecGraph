import os
import docx

doc_path = os.path.abspath('SecGraph_Project_Report.docx')
doc = docx.Document(doc_path)

print(f"File: {doc_path}")
print(f"Size: {os.path.getsize(doc_path):,} bytes")
print(f"Paragraphs: {len(doc.paragraphs)}")
print(f"Tables: {len(doc.tables)}")
print(f"Sections: {len(doc.sections)}")
print(f"Inline Pictures: {len(doc.inline_shapes)}")

text_len = sum(len(p.text) for p in doc.paragraphs)
table_text_len = sum(len(c.text) for t in doc.tables for row in t.rows for c in row.cells)
total_words = (text_len + table_text_len) // 6

print(f"Total Text Characters: {text_len + table_text_len:,}")
print(f"Estimated Word Count: {total_words:,} words")
